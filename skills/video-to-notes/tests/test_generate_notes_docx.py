import os
import sys
import tempfile
from pathlib import Path
import docx

# 將 scripts 加入模組搜尋路徑
script_dir = Path(__file__).resolve().parent.parent / "scripts"
sys.path.insert(0, str(script_dir))

from generate_notes_docx import (
    MeetingNotesParser,
    NotesDocxBuilder,
    MeetingNotesData,
)


SAMPLE_MEETING_MD = """---
title : 2026年度雙軌品質管理專案進度與規格同步會
description : 檢視 Q3 導流板試產良率、決議第 2 套治具採購時程與後續行動待辦。
date : 2026-09-14 14:00
aliases : []
status: unread
tags : 
- Type/📝/✨
Topics : 專案會議
---

# 2026年度雙軌品質管理專案進度與規格同步會

## 1. 會議基本資訊
- 會議時間: 2026-09-14 14:00 - 15:30
- 會議地點: 研發部第 3 會議室 (Teams 線上同步)
- 會議主席: 張處長
- 會議記錄: 林助理
- 出席人員: 張處長、王組長、李工程師、陳品質專員
- 請假人員: 無

## 2. 會議核心摘要 (Highlights)
- Q3 導流板試產公差已收斂至 0.05mm 以內，整體良率顯著提升至 99.2%。
- 針對新導入之自動化設備，維護工程師已完成第一階段操作影片與 SOP 撰寫。
- 確認產線目前產能瓶頸主要在第 1 套治具之換模調機時間。

## 3. 關鍵決策事項 (Decisions Made)
| 編號 | 決策主題 | 決策內容與共識 | 提案人 / 負責人 | 生效日期 |
|---|---|---|---|---|
| D-01 | 治具擴充採購 | 同意採購第 2 套導流板治具，預算控制在 15 萬元內 | 王組長 | 2026-09-20 |
| D-02 | SOP 審查排程 | 全廠保養 SOP 統一於每季末由品管與現場共同覆核 | 張處長 | 2026-10-01 |

## 4. 待辦事項清單 (Action Items / Todo List)
| 任務項目 | 負責人 | 截止期限 | 當前狀態 |
|---|---|---|---|
| 完成第 2 套治具供應商詢價與規格確認 | 王組長 | 2026-09-18 | 進行中 |
| 統整研磨機台保養逐字稿並完成錯別字校對 | 李工程師 | 2026-09-22 | 未開始 |
| 安排品管小組進行治具公差盲測驗證 | 陳品質專員 | 2026-09-25 | 未開始 |

## 5. 各議題討論紀要 (Agenda & Discussions)

### 議題一: 導流板試產公差與良率檢討
- 【張處長】簡報中顯示目前良率達 99.2%，但邊緣毛邊仍有零星異常，品管端需加強抽檢頻率。
- 【李工程師】現場換刀塊後已改善切削穩定度，預計下週可完全消除微細毛邊。

### 議題二: 換模換線調機效率改善
- 【王組長】目前調機耗時過長主因在於手動鎖固對位，若引進第 2 套治具可平行預裝模具，預估節省 40% 工時。
- 【張處長】請王組長儘速送出採購請購單。

## 6. 下次會議追蹤項目 (Next Meeting Follow-ups)
| 追蹤項目 | 預計報告人 | 期望產出 / 查核標準 |
|---|---|---|
| 第 2 套治具採購進度與交期回報 | 王組長 | 供應商報價單與交期確認回簽 |
| 導流板零毛邊試產驗證數據 | 李工程師 | 50 件樣品量測數據報告 (Cpk > 1.33) |

# 參考資料
- [專案執行/視訊會議/20260914_雙軌品質專案會議.mp4]
"""


def test_meeting_markdown_parser():
    """驗證 Markdown 會議記錄文字解析器"""
    data = MeetingNotesParser.parse(SAMPLE_MEETING_MD)

    # 1. 驗證標題與 YAML frontmatter
    assert "2026年度雙軌品質管理專案進度與規格同步會" in data.title
    assert "status: unread" in data.parameters_yaml
    assert "Type/📝/✨" in data.parameters_yaml

    # 2. 驗證基本資訊
    assert data.chairperson == "張處長"
    assert data.minute_taker == "林助理"
    assert "李工程師" in data.attendees

    # 3. 驗證核心摘要
    assert len(data.highlights) >= 3
    assert any("99.2%" in h for h in data.highlights)

    # 4. 驗證決策事項
    assert len(data.decisions) == 2
    assert data.decisions[0].id == "D-01"
    assert "治具擴充採購" in data.decisions[0].topic
    assert "王組長" in data.decisions[0].owner_or_proposer

    # 5. 驗證 Todo 待辦事項
    assert len(data.todos) == 3
    assert "王組長" in data.todos[0].owner
    assert "2026-09-18" in data.todos[0].due_date

    # 6. 驗證發言人討論紀要
    assert len(data.agenda_discussions) == 2
    assert data.agenda_discussions[0].topic_title == "導流板試產公差與良率檢討"
    assert len(data.agenda_discussions[0].discussions) >= 2
    assert data.agenda_discussions[0].discussions[0].speaker_name == "張處長"

    # 7. 驗證下次會議追蹤項目
    assert len(data.next_meeting_followups) == 2
    assert "第 2 套治具採購進度" in data.next_meeting_followups[0].item_title
    assert data.next_meeting_followups[0].responsible_person == "王組長"

    # 8. 驗證文末參考資料
    assert len(data.references) >= 1
    assert "20260914_雙軌品質專案會議.mp4" in data.references[0]


def test_zero_screenshot_docx_builder():
    """驗證 Word 生成器具備 0 截圖防呆保證且排版完整"""
    data = MeetingNotesParser.parse(SAMPLE_MEETING_MD)
    builder = NotesDocxBuilder(data)
    doc = builder.build()

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        doc.save(tmp_path)
        loaded_doc = docx.Document(tmp_path)

        # 1. 關鍵驗證：全文件絕無圖片嵌入 (Zero-screenshot Guarantee)
        assert len(loaded_doc.inline_shapes) == 0

        # 2. 驗證表格數量 (基本資訊、決策、Todo、下次會議追蹤、YAML)
        assert len(loaded_doc.tables) >= 4

        # 3. 驗證 Todo 表格包含核取方塊符號 ☐
        todo_table_found = False
        for table in loaded_doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if "☐" in cell.text:
                        todo_table_found = True
                        break
        assert todo_table_found, "Todo 表格中未找到核取方塊符號 ☐"

        # 4. 驗證下次會議追蹤項目被正確輸出
        followup_found = False
        for p in loaded_doc.paragraphs:
            if "下次會議追蹤項目" in p.text:
                followup_found = True
                break
        assert followup_found, "Word 文件中未找到下次會議追蹤項目章節標題"

    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_edge_cases_empty_decisions_and_todos():
    """驗證缺少決策事項或 Todo 時的優雅降級容錯能力"""
    minimal_md = """---
title : 臨時現況快報會
description : 
date : 2026-09-14
aliases : []
status: unread
tags : 
- Type/📝/✨
Topics : 
---

# 臨時現況快報會

## 2. 會議核心摘要 (Highlights)
- 各產線稼動率均維持在 85% 以上。

## 5. 各議題討論紀要 (Agenda & Discussions)
- 【主管】目前無重大異常回報，請各組依規劃進行。

# 參考資料
- [臨時會錄音.mp3]
"""
    data = MeetingNotesParser.parse(minimal_md)
    assert len(data.decisions) == 0
    assert len(data.todos) == 0
    assert len(data.next_meeting_followups) == 0

    builder = NotesDocxBuilder(data)
    doc = builder.build()

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        doc.save(tmp_path)
        loaded = docx.Document(tmp_path)
        # 即使無決策與 Todo，Word 文件仍可正常建立且無圖
        assert len(loaded.inline_shapes) == 0
        assert any("各產線稼動率" in p.text for p in loaded.paragraphs)
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_speaker_attribution_and_varied_formats():
    """驗證多樣發言人標籤格式與綜合討論解析"""
    custom_md = """---
title : 研發部週會
date : 2026-09-15
---

# 研發部週會

## 5. 各議題討論紀要
### 議題 A
- 【王組長 / 機械工程】回報夾爪公差已符合客戶需求。
- [李專員] 測試數據報告將於明天送出。
- 現場同仁共同確認機構干涉問題已排除。

# 參考資料
- [週會.mp4]
"""
    data = MeetingNotesParser.parse(custom_md)
    assert len(data.agenda_discussions) == 1
    discussions = data.agenda_discussions[0].discussions
    assert len(discussions) == 3
    assert discussions[0].speaker_name == "王組長 / 機械工程"
    assert "客戶需求" in discussions[0].statement_summary
    assert discussions[2].speaker_name == "綜合討論"


def test_cleanup_temp_safety_validator():
    """驗證暫存清理腳本防呆機制"""
    from cleanup_temp import is_safe_temp_path
    
    # 正常合法路徑
    safe_path = Path(r"C:\Users\user\.gemini\antigravity-ide\brain\123456\.tempmediaStorage")
    assert is_safe_temp_path(safe_path) is True

    # 非法路徑（未包含 .tempmediaStorage 或太短）
    assert is_safe_temp_path(Path(r"C:\Users\user\Desktop")) is False
    assert is_safe_temp_path(Path(r"C:\Windows\System32")) is False
    assert is_safe_temp_path(Path(r"C:\project\.tempmediaStorage")) is False

